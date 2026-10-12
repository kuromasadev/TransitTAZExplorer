from pathlib import Path
from openpyxl import load_workbook
import json,collections,sys,xml.etree.ElementTree as ET
sys.stdout.reconfigure(encoding='utf-8')
r=Path(__file__).resolve().parents[1]; out=r/'project-review'
inv=json.loads((out/'inventory.json').read_text(encoding='utf-8'))
for d in inv:
 if 'property_keys' in d:print(d['path'],d['feature_count'],d['property_keys'], 'nulls', {k:v for k,v in d['null_counts'].items() if v})
 if d['path'].endswith('.csv'):print(d['path'],d['rows'],d['headers'])
wb=load_workbook(r/'RAW/pnr_esrigc/TAZ DATA.xlsx',read_only=True,data_only=True)
tables={s.title:[dict(zip(next(s.values),row)) for row in list(s.values)[1:]] for s in wb}
print('TAZ uniqueness',len(set(x['TAZ *'] for x in tables['TAZ'])))
print('SECT sectors',sorted(set(x['Sector'] for x in tables['SECT'])))
print('RELATE join_count distribution',dict(collections.Counter(x['Join_Count'] for x in tables['relate1'])))
odsector=collections.defaultdict(set); sectors={x['TAZ']:x['Sector'] for x in tables['SECT']}
for x in tables['TAZ']:odsector[x['O-D AREAS *']].add(sectors.get(x['TAZ *']))
print('MULTI SECTOR OD', {k:sorted(v) for k,v in odsector.items() if len(v)>1})
print('OD workbook names vs TAZ names',len(set(odsector)),set(x['ODAREA'] for x in tables['OD'])-set(odsector))
ca=load_workbook(r/'RAW/pnr_esrigc/community_analyst_export.xlsx',read_only=True,data_only=True)
for s in ca:
 rows=list(s.values);ids=set(x[-1] for x in rows[1:]);print('CA',s.title,'IDs',sorted(ids),'duplicate keys',len(rows[1:])-len(set((x[0],x[-1]) for x in rows[1:])))
pnr=load_workbook(r/'RAW/pnr_plnfmm/ALL PNR.xlsx',data_only=False)
for s in pnr:
 print('PNR TABLES',s.title,[(t.name,t.ref) for t in s.tables.values()])
 for t in s.tables.values():
  if t.name=='ApproxRelate':
   cells=list(s[t.ref]); data=[[c.value for c in row] for row in cells];print('APPROX RELATE',len(data)-1,data[:15])
for name in ['15','3','100']:
 s=load_workbook(r/f'RAW/ltsv_esrica/{name}.xlsx',read_only=True,data_only=True).active
 print('REPORT',name)
 for i,row in enumerate(s.values):
  if i in [0,1,2,3,48,49,50,51,52,53,54,55,138,139,140,141,142,143,144,145,146,147,148,149]:print(i+1,row)
for p in (r/'DB/.dbeaver').glob('*json*'):
 j=json.loads(p.read_text());print('DB',p.name,'keys',list(j))
 for c in j.get('connections',{}).values():print('connection provider',c.get('provider'),'driver',c.get('driver'),'configuration keys',list(c.get('configuration',{})))
for p in (r/'BI').glob('*.twb'):
 x=ET.parse(p).getroot(); print('TABLEAU',p.name,'relationship count',len(x.findall('.//relationship')),'relationships', [ET.tostring(t,encoding='unicode') for t in x.findall('.//relationship')][:20])
